import logging
import uuid
from functools import lru_cache

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.db.models import Prefetch, Q
from django.utils.encoding import force_str
from django.utils.http import urlsafe_base64_decode
from rest_framework.exceptions import ValidationError

from accounts.emails import account_activation_token
from accounts.invite_store import (
    delete_invite_token_pair,
    get_user_id_for_token,
    revoke_invite_tokens_for_user,
    store_invite_token,
)
from accounts.models import Profile, Role, User
from core.exceptions import AppNotFound
from utils.files import django_field_file_sha256, file_content_hash

logger = logging.getLogger(__name__)


def get_user_by_pk(*, pk, app_not_found_code=None):
    """
    Return a user by primary key.

    If the user does not exist: raises ``AppNotFound("UserNotFound")`` by default, or
    ``AppNotFound(app_not_found_code)`` when ``app_not_found_code`` is set (PascalCase i18n key).
    """
    try:
        return User.objects.get(pk=pk)
    except User.DoesNotExist:
        raise AppNotFound(app_not_found_code or "UserNotFound")


def list_users():
    """
    Return all active (non–soft-deleted) users with roles prefetched for listing.
    """
    return (
        User.objects.all()
        .prefetch_related(
            Prefetch(
                "groups",
                queryset=Group.objects.select_related("role"),
            )
        )
        .order_by("email")
    )


def get_user_with_roles_prefetched(*, pk):
    """
    Return a non-deleted user by pk with ``groups`` + ``role`` prefetched
    (same shape as :func:`list_users` rows).
    """
    try:
        return (
            User.objects.all()
            .prefetch_related(
                Prefetch(
                    "groups",
                    queryset=Group.objects.select_related("role"),
                )
            )
            .get(pk=pk)
        )
    except User.DoesNotExist:
        raise AppNotFound("UserNotFound")


@lru_cache(maxsize=1)
def _can_teach_permission():
    return Permission.objects.get(
        content_type__app_label="accounts",
        codename="can_teach",
    )


def _teaching_users_queryset():
    perm = _can_teach_permission()
    return User.objects.filter(
        Q(groups__permissions=perm) | Q(user_permissions=perm)
    ).distinct()


def list_teachers():
    """
    Return users who have ``accounts.can_teach`` (via group or direct permission).
    """
    return _teaching_users_queryset().order_by("email")


def user_can_teach(*, user) -> bool:
    """
    True if ``user`` has ``accounts.can_teach`` (via group or direct permission).
    Same membership rule as :func:`list_teachers`.
    """
    return _teaching_users_queryset().filter(pk=user.pk).exists()


def get_teacher_user_by_pk(*, pk):
    """
    Return a user by primary key only if they have ``accounts.can_teach``.

    Raises ``AppNotFound("TeacherNotFound")`` if the user does not exist,
    ``ValidationError`` with ``UserCannotTeach`` if they exist but cannot teach.
    Uses at most two database queries (one when the user is a valid teacher).
    """
    perm = _can_teach_permission()
    teacher = (
        User.objects.filter(pk=pk)
        .filter(Q(groups__permissions=perm) | Q(user_permissions=perm))
        .distinct()
        .first()
    )
    if teacher is not None:
        return teacher
    if User.objects.filter(pk=pk).exists():
        raise ValidationError({"messages": ["UserCannotTeach"]})
    raise AppNotFound("TeacherNotFound")


def assert_user_can_teach(*, user):
    """Raise ``ValidationError`` if the user is not allowed to act as a teacher."""
    if not user_can_teach(user=user):
        raise ValidationError({"messages": ["UserCannotTeach"]})


@lru_cache(maxsize=1)
def _can_participate_in_coaching_permission():
    return Permission.objects.get(
        content_type__app_label="accounts",
        codename="can_participate_in_coaching",
    )


def _coaching_participant_users_queryset():
    perm = _can_participate_in_coaching_permission()
    return User.objects.filter(
        Q(groups__permissions=perm) | Q(user_permissions=perm)
    ).distinct()


def user_can_participate_in_coaching(*, user) -> bool:
    """
    True if ``user`` has ``accounts.can_participate_in_coaching`` (via group or direct).
    Same membership rule as roster eligibility checks in coachings / classes.
    """
    return _coaching_participant_users_queryset().filter(pk=user.pk).exists()


def assert_user_can_participate_in_coaching(*, user):
    """Raise ``ValidationError`` with ``UserCannotParticipateInCoaching`` if not eligible."""
    if not user_can_participate_in_coaching(user=user):
        raise ValidationError({"messages": ["UserCannotParticipateInCoaching"]})


def assert_user_ids_can_participate_in_coaching(*, user_ids):
    """
    Raise ``ValidationError`` with ``UserCannotParticipateInCoaching`` if any listed user
    lacks ``accounts.can_participate_in_coaching`` or does not exist.

    Used when bulk-assigning coaching participants so roster inserts cannot bypass
    permission checks (e.g. classe members created before enforcement).
    """
    unique_ids = list(dict.fromkeys(user_ids))
    if not unique_ids:
        return
    perm = _can_participate_in_coaching_permission()
    allowed = set(
        User.objects.filter(pk__in=unique_ids)
        .filter(Q(groups__permissions=perm) | Q(user_permissions=perm))
        .distinct()
        .values_list("pk", flat=True)
    )
    if len(allowed) != len(set(unique_ids)):
        raise ValidationError({"messages": ["UserCannotParticipateInCoaching"]})


@transaction.atomic
def create_user(
    *,
    email,
    first_name,
    last_name,
    role_code,
    created_by,
    phone="",
    phone_country_code="",
):
    """
    Create a new user with a role assignment.
    """
    try:
        role = Role.objects.select_related("group").get(code=role_code)
    except Role.DoesNotExist:
        raise ValidationError({"messages": ["RoleNotFound"]})

    try:
        user = User.objects.create_user(
            email=email,
            password=None,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            phone_country_code=phone_country_code,
            is_active=False,
            created_by=created_by,
        )
    except ValueError:
        # UserManager.create_user raises ValueError for falsy email — not a DRF
        # exception, so it would bypass custom_exception_handler and become a 500.
        raise ValidationError({"messages": ["EmailIsMandatory"]})

    user.groups.add(role.group)

    return user


def get_profile_for_user(*, user):
    try:
        return Profile.objects.get(user=user)
    except Profile.DoesNotExist:
        raise AppNotFound("ProfileNotFound")


def create_profile_for_user(*, user, created_by=None, **fields):
    profile = Profile(user=user, created_by=created_by, **fields)
    try:
        profile.full_clean()
    except DjangoValidationError as exc:
        errs = getattr(exc, "error_dict", None) or {}
        if "user" in errs:
            raise ValidationError({"messages": ["ProfileAlreadyExists"]}) from exc
        raise
    try:
        profile.save()
    except IntegrityError:
        raise ValidationError({"messages": ["ProfileAlreadyExists"]})
    return profile


def update_profile_for_user(*, profile, updated_by=None, **fields):
    new_photo = fields.get("photo")
    if new_photo is not None and profile.photo:
        if file_content_hash(new_photo) == django_field_file_sha256(profile.photo):
            fields = {k: v for k, v in fields.items() if k != "photo"}

    for key, value in fields.items():
        setattr(profile, key, value)
    profile.updated_by = updated_by
    profile.full_clean()
    profile.save()
    return profile


@transaction.atomic
def update_user_core_fields(*, pk, updated_by, **updates):
    """
    Partial update of principal user fields (name, contact, language, parent, address).

    Intended for callers that hold ``accounts.change_user`` (direction).
    """
    if not updates:
        return get_user_with_roles_prefetched(pk=pk)

    user = get_user_by_pk(pk=pk)
    for key, value in updates.items():
        setattr(user, key, value)
    user.updated_by = updated_by
    try:
        user.full_clean()
        user.save()
    except DjangoValidationError as exc:
        error_dict = getattr(exc, "error_dict", {}) or {}
        if "email" in error_dict:
            raise ValidationError({"messages": ["EmailAlreadyExists"]}) from exc
        raise ValidationError({"messages": ["InvalidUserData"]}) from exc
    except IntegrityError as exc:
        raise ValidationError({"messages": ["EmailAlreadyExists"]}) from exc

    logger.info(
        "User core fields updated: pk=%s updated_by=%s fields=%s",
        pk,
        getattr(updated_by, "email", updated_by),
        sorted(updates.keys()),
    )
    return get_user_with_roles_prefetched(pk=user.pk)


def activate_account(*, uid, token, password):
    User = get_user_model()
    try:
        pk = force_str(urlsafe_base64_decode(uid))
        account_user = User.objects.get(pk=pk)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        raise ValidationError({"messages": ["InvalidActivationLink"]})

    if account_user.is_active:
        raise ValidationError({"messages": ["AccountAlreadyActive"]})

    if not account_activation_token.check_token(account_user, token):
        raise ValidationError({"messages": ["InvalidOrExpiredToken"]})

    account_user.set_password(password)
    account_user.is_active = True
    account_user.save()

    logger.info("Account activated: %s", account_user.email)
    return account_user


# --- Invites & onboarding ----------------------------------------------------

INVITE_TARGETS_DIRECTION = frozenset(
    {"coordination", "collaborator", "teacher", "parent", "student"}
)
INVITE_TARGETS_COORDINATION = frozenset({"collaborator", "teacher", "parent", "student"})


def _invite_scope_for_user(*, inviter) -> str | None:
    codes = set()
    for group in inviter.groups.select_related("role").all():
        role = getattr(group, "role", None)
        if role:
            codes.add(role.code)
    if "direction" in codes:
        return "direction"
    if "coordination" in codes:
        return "coordination"
    return None


def assert_inviter_can_assign_role(*, inviter, target_role_code: str) -> None:
    """
    Enforce direction/coordination hierarchy for invite targets.
    """
    try:
        Role.objects.get(code=target_role_code)
    except Role.DoesNotExist:
        raise ValidationError({"messages": ["RoleNotFound"]})

    scope = _invite_scope_for_user(inviter=inviter)
    if scope == "direction":
        allowed = INVITE_TARGETS_DIRECTION
    elif scope == "coordination":
        allowed = INVITE_TARGETS_COORDINATION
    else:
        raise ValidationError({"messages": ["CannotInviteRole"]})

    if target_role_code not in allowed:
        raise ValidationError({"messages": ["CannotInviteRole"]})


def get_user_primary_role_code(*, user) -> str | None:
    for group in user.groups.select_related("role").all():
        role = getattr(group, "role", None)
        if role:
            return role.code
    return None


def _enqueue_invite_email(*, user_id: int, token: str) -> None:
    from accounts.tasks import send_invite_email_task

    send_invite_email_task.delay(user_id, token)


@transaction.atomic
def invite_user(*, email, role_code, created_by):
    """
    Create an inactive user with the given role and issue an onboarding token (Redis + email task).
    """
    email_n = User.objects.normalize_email(email)
    assert_inviter_can_assign_role(inviter=created_by, target_role_code=role_code)

    if User.all_objects.filter(email__iexact=email_n).exists():
        raise ValidationError({"messages": ["EmailAlreadyExists"]})

    role = Role.objects.select_related("group").get(code=role_code)

    try:
        user = User.objects.create_user(
            email=email_n,
            password=None,
            first_name="",
            last_name="",
            phone="",
            phone_country_code="",
            is_active=False,
            created_by=created_by,
        )
    except ValueError:
        raise ValidationError({"messages": ["EmailIsMandatory"]}) from None

    user.groups.add(role.group)

    token = str(uuid.uuid4())

    def _after_commit(uid=user.pk, t=token):
        store_invite_token(user_id=uid, token=t)
        _enqueue_invite_email(user_id=uid, token=t)

    transaction.on_commit(_after_commit)

    logger.info(
        "User invite issued: email=%s role=%s invited_by=%s",
        email_n,
        role_code,
        getattr(created_by, "email", created_by),
    )
    return get_user_with_roles_prefetched(pk=user.pk)


def _parse_bulk_invite_headers(ws):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValidationError({"messages": ["BulkInviteFileInvalid"]})
    headers = [
        (str(h).strip().lower() if h is not None else "") for h in header_row
    ]
    try:
        email_idx = headers.index("email")
        role_idx = headers.index("role")
    except ValueError as exc:
        raise ValidationError({"messages": ["BulkInviteFileInvalid"]}) from exc
    return email_idx, role_idx


def parse_bulk_invite_xlsx_rows(*, uploaded_file):
    """
    Read an .xlsx with ``email`` and ``role`` columns (header row required).
    Returns a list of dicts: ``row`` (1-based sheet row), ``email``, ``role`` (raw strings).
    """
    from openpyxl import load_workbook

    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    try:
        wb = load_workbook(uploaded_file, read_only=False, data_only=True)
    except Exception as exc:
        raise ValidationError({"messages": ["BulkInviteFileInvalid"]}) from exc

    max_data_row = 5002
    parsed = []
    try:
        ws = wb.active
        email_idx, role_idx = _parse_bulk_invite_headers(ws)
        for row_idx, row_tuple in enumerate(
            ws.iter_rows(min_row=2, max_row=max_data_row, values_only=True),
            start=2,
        ):
            if row_tuple is None:
                continue
            vals = list(row_tuple)
            if not any(v not in (None, "") for v in vals):
                continue
            email_cell = vals[email_idx] if email_idx < len(vals) else None
            role_cell = vals[role_idx] if role_idx < len(vals) else None
            parsed.append(
                {
                    "row": row_idx,
                    "email": str(email_cell).strip() if email_cell is not None else "",
                    "role": str(role_cell).strip() if role_cell is not None else "",
                }
            )
    finally:
        wb.close()

    return parsed


def _validate_bulk_invite_rows(*, inviter, parsed_rows):
    """
    All-or-nothing validation. Returns normalized rows:
    ``[{"email": normalized_email, "role_code": str}, ...]``.
    """
    row_errors = []
    seen_emails = set()
    valid_codes = {r.code for r in Role.objects.only("code")}

    # Compute inviter scope once — avoids N Role DB lookups inside the loop.
    scope = _invite_scope_for_user(inviter=inviter)
    if scope == "direction":
        allowed_roles = INVITE_TARGETS_DIRECTION
    elif scope == "coordination":
        allowed_roles = INVITE_TARGETS_COORDINATION
    else:
        allowed_roles = frozenset()

    # Batch-check email existence — avoids N queries inside the loop.
    candidate_emails = [
        User.objects.normalize_email(e["email"]).lower()
        for e in parsed_rows
        if e["email"]
    ]
    existing_emails = set(
        User.all_objects.filter(email__in=candidate_emails).values_list("email", flat=True)
    )

    for entry in parsed_rows:
        row_num = entry["row"]
        raw_email = entry["email"]
        raw_role = entry["role"]
        errs = []

        if not raw_email:
            errs.append("InvalidEmail")
        else:
            try:
                validate_email(raw_email)
            except DjangoValidationError:
                errs.append("InvalidEmail")

        email_norm = User.objects.normalize_email(raw_email).lower() if raw_email else ""

        role_code = raw_role.strip().lower() if raw_role else ""
        if not role_code:
            errs.append("RoleNotFound")
        elif role_code not in valid_codes:
            errs.append("RoleNotFound")

        if email_norm and email_norm in seen_emails:
            errs.append("DuplicateEmailInFile")
        elif email_norm:
            seen_emails.add(email_norm)

        if email_norm and email_norm in existing_emails:
            errs.append("EmailAlreadyExists")

        if not errs and role_code:
            if role_code not in allowed_roles:
                errs.append("CannotInviteRole")

        if errs:
            row_errors.append(
                {
                    "row": row_num,
                    "email": raw_email or "",
                    "errors": errs,
                }
            )

    if row_errors:
        raise ValidationError(
            {"messages": ["BulkInviteValidationFailed"], "rows": row_errors}
        )

    return [
        {
            "email": User.objects.normalize_email(e["email"]).lower(),
            "role_code": e["role"].strip().lower(),
        }
        for e in parsed_rows
    ]


def bulk_invite_users(*, inviter, uploaded_file):
    """
    Parse and validate an Excel file; create all users atomically; enqueue emails after commit.
    """
    parsed = parse_bulk_invite_xlsx_rows(uploaded_file=uploaded_file)
    if not parsed:
        raise ValidationError({"messages": ["BulkInviteFileInvalid"]})

    normalized_rows = _validate_bulk_invite_rows(inviter=inviter, parsed_rows=parsed)

    users_out = []
    with transaction.atomic():
        role_codes = {row["role_code"] for row in normalized_rows}
        roles_by_code = {
            r.code: r
            for r in Role.objects.select_related("group").filter(code__in=role_codes)
        }
        for row in normalized_rows:
            role = roles_by_code[row["role_code"]]
            try:
                user = User.objects.create_user(
                    email=row["email"],
                    password=None,
                    first_name="",
                    last_name="",
                    phone="",
                    phone_country_code="",
                    is_active=False,
                    created_by=inviter,
                )
            except ValueError:
                raise ValidationError({"messages": ["EmailIsMandatory"]}) from None
            user.groups.add(role.group)
            users_out.append(user)

        def _after_commit():
            for u in users_out:
                tok = str(uuid.uuid4())
                store_invite_token(user_id=u.pk, token=tok)
                _enqueue_invite_email(user_id=u.pk, token=tok)

        transaction.on_commit(_after_commit)

    logger.info(
        "Bulk invite issued: count=%s invited_by=%s",
        len(users_out),
        getattr(inviter, "email", inviter),
    )
    return [get_user_with_roles_prefetched(pk=u.pk) for u in users_out]


def get_onboarding_preview(*, token: str):
    """
    Resolve invite token to email + role code for the public GET onboard endpoint.
    """
    uid = get_user_id_for_token(token=token)
    if uid is None:
        raise AppNotFound("InvalidOrExpiredInviteToken")

    try:
        user = User.objects.get(pk=uid)
    except User.DoesNotExist:
        raise AppNotFound("InvalidOrExpiredInviteToken")

    role_code = get_user_primary_role_code(user=user)
    if role_code is None:
        raise AppNotFound("InvalidOrExpiredInviteToken")

    return {"email": user.email, "role": role_code}


@transaction.atomic
def complete_onboarding(
    *,
    token,
    first_name,
    last_name,
    password,
    phone,
    phone_country_code,
    date_of_birth,
):
    uid = get_user_id_for_token(token=token)
    if uid is None:
        raise ValidationError({"messages": ["InvalidOrExpiredInviteToken"]})

    try:
        user = User.objects.get(pk=uid)
    except User.DoesNotExist:
        raise ValidationError({"messages": ["InvalidOrExpiredInviteToken"]})

    if user.is_active:
        raise ValidationError({"messages": ["AccountAlreadyActive"]})

    user.first_name = first_name
    user.last_name = last_name
    user.phone = phone
    user.phone_country_code = phone_country_code
    user.date_of_birth = date_of_birth
    user.set_password(password)
    user.is_active = True
    user.updated_by = user
    user.full_clean()
    user.save()

    tok = token

    def _cleanup(u=user.pk, t=tok):
        delete_invite_token_pair(user_id=u, token=t)

    transaction.on_commit(_cleanup)

    logger.info("Onboarding completed: user_id=%s email=%s", user.pk, user.email)

    from rest_framework_simplejwt.tokens import RefreshToken

    refresh = RefreshToken.for_user(user)
    return {
        "access": str(refresh.access_token),
        "refresh": str(refresh),
    }


def resend_invite(*, user_pk, requesting_user):
    """
    Invalidate any existing invite token and issue a new one for an inactive user.
    No DB writes — Redis ops and email dispatch run directly (no on_commit needed).
    """
    user = get_user_by_pk(pk=user_pk)

    if user.is_active:
        raise ValidationError({"messages": ["AccountAlreadyActive"]})

    role_code = get_user_primary_role_code(user=user)
    if role_code is None:
        raise ValidationError({"messages": ["CannotInviteRole"]})

    assert_inviter_can_assign_role(inviter=requesting_user, target_role_code=role_code)

    new_token = str(uuid.uuid4())
    revoke_invite_tokens_for_user(user_id=user.pk)
    store_invite_token(user_id=user.pk, token=new_token)
    _enqueue_invite_email(user_id=user.pk, token=new_token)

    logger.info(
        "Invite resent: user_id=%s requested_by=%s",
        user.pk,
        getattr(requesting_user, "email", requesting_user),
    )
